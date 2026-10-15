"""Bouw de Nederlandse ParkHaven-presentatie v0.2 naast de bewaarde v0.1."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_progress_presentation import (
    BG,
    BLUE,
    CAPTURE_ROOT,
    GOLD,
    GREEN,
    H,
    INK,
    LINE,
    LOGO,
    MID_GREEN,
    MINT,
    MUTED,
    ORANGE,
    PAPER,
    PURPLE,
    ROOT,
    SOFT_BLUE,
    TEAL,
    W,
    WHITE,
    add_notes,
    arrow,
    metric_card,
    rect,
    rich_text,
    step_card,
    text,
    title,
)


OUTPUT_DIR = ROOT / "presentations"
OUTPUT = OUTPUT_DIR / "ParkHaven-presentatie-v0.2.pptx"
SCRIPT_OUTPUT = OUTPUT_DIR / "ParkHaven-presentatiescript-v0.2.md"


scripts = [
    """ParkHaven Studio is een lokale ontwerpomgeving die een Rhino-model omzet in een begrijpelijke ontwerpbeoordeling. De app brengt Rhino, Grasshopper, modelbeelden, meetresultaten en rapportage samen in één begeleide route. De gebruiker importeert een ontwerp, plaatst het op de locatie, analyseert het en maakt een controleerbaar verslag. Eerst bekijken we die hoofdwerking; daarna zoomen we in op de technische opbouw.""",
    """Het huidige proces is vaak versnipperd. Geometrie staat in Rhino, berekeningen in Grasshopper, beelden worden handmatig gemaakt en conclusies belanden in losse documenten. Daardoor is moeilijk te herleiden welke uitkomst bij welke modelversie hoort. ParkHaven Studio verbindt die onderdelen, bewaart de bron en koppelt iedere analyse, afbeelding en conclusie aan dezelfde modelstatus. Zo wordt itereren overzichtelijker en beter uitlegbaar.""",
    """Voor de gebruiker zijn er vijf stappen. Eerst wordt een 3DM-bestand geïmporteerd. Daarna controleert de ontwerper positie, rotatie en modeleenheid met een snelle 3D-preview. Vervolgens start de lokale Rhino- en Grasshopper-analyse. Na afloop toont de app prestaties en vier vaste modelbeelden. Optioneel volgt begrensd AI-advies. Tot slot ontstaat een leesbaar rapport dat als PDF kan worden opgeslagen.""",
    """De centrale werkruimte is de 3D-viewer. Hier staat het ontwerp naast gevels, parkdoel, referentiegebouwen en de lokale luchtfoto. Lagen kunnen afzonderlijk aan en uit, en vaste camera's maken varianten vergelijkbaar. Ontwerpvolumes houden overal dezelfde kleuren: blauw, oranje en paars. Zo blijven Volume A, B en C herkenbaar in de viewer, de bewijsbeelden en het rapport.""",
    """De app levert kerngetallen zoals bruto vloeroppervlak, zicht en indicatieve kosten. Daarnaast zijn er vergelijkende screenings voor zon, thermisch comfort en fundering; deze zijn geen definitief ingenieursadvies. Alles komt samen in een ontwerpverslag met modelbeelden, meetresultaten, aannames en eventuele ontwerpacties. Het resultaat is dus niet alleen een dashboard, maar een dossier dat de onderbouwing van een ontwerpkeuze zichtbaar maakt.""",
    """Onder de interface werken zes onderdelen samen. Een Windows-programma met WebView2 toont de app. De lokale FastAPI-backend beheert projecten, runs en rapporten. Three.js verzorgt de 3D-viewer. Voor analyse opent de backend Rhino zichtbaar; een eigen plug-in laadt de Grasshopper-definitie en exporteert JSON-resultaten. Een optionele AI-provider ontvangt alleen begrensde tekst. Beelden blijven lokaal en het rapport wordt lokaal opgebouwd.""",
    """Iedere analyse is technisch een geïsoleerde run. Het originele 3DM-bestand blijft ongewijzigd. Per run maakt de app een modelkopie, jobbestand, Grasshopper-definitie, weerbestand en datamap. Meetwaarden, scene en geometriesamenvatting worden atomair gepubliceerd. SHA-256-hashes koppelen bron, definitie, weerdata, resultaten en beelden. Daardoor kan de app verouderd of verwisseld bewijs weigeren en overschrijft een nieuwe analyse niet ongemerkt een eerder resultaat.""",
    """De bewijs- en AI-keten heeft een harde datagrens. De viewer maakt vier gestandaardiseerde beelden met bekende camera's en lagen. De backend bouwt daarnaast een compacte tekstbeschrijving van volumes en meetwaarden. Alleen die tekst mag naar een AI-endpoint; afbeeldingen en ruwe geometrie blijven lokaal. Het antwoord wordt gevalideerd op bekende volumes en meetwaarden. Verschuivingen zijn toegestaan als verkennende tests, maar zonder geverifieerde plot- of tunnelcontour gelden ze nooit als bewijs van haalbaarheid.""",
    """De basisworkflow werkt nu van import tot PDF en is getest met het echte Parkhaven-model. De huidige versie ondersteunt projectopslag, plaatsing, Rhino- en Grasshopper-analyse, 3D-weergave, bewijsbeelden, begrensd AI-advies en rapportage. De volgende fase richt zich op betrouwbaarheid: geverifieerde locatiegrenzen, gekalibreerde normen, ondertekende distributie en tests op schone machines. De kern staat; v0.2 bereidt de app voor op bredere praktijktesten.""",
]


def footer_nl(slide) -> None:
    slide.shapes.add_picture(str(LOGO), Inches(0.62), Inches(6.96), height=Inches(0.37))
    rich_text(slide, [("PARKHAVEN ", INK, True), ("STUDIO", MUTED, False)],
              1.02, 7.03, 2.2, 0.2, 8)
    text(slide, "LOKAAL • HERLEIDBAAR • MODELGEBASEERD", 8.75, 7.03, 3.95, 0.2,
         7, MUTED, True, align=PP_ALIGN.RIGHT)


def pill(slide, label, x, y, w, color=GREEN, fill=PAPER) -> None:
    rect(slide, x, y, w, 0.38, fill, color, True)
    text(slide, label, x + 0.08, y + 0.105, w - 0.16, 0.16, 8, color, True,
         align=PP_ALIGN.CENTER)


def component(slide, x, y, w, h, label, detail, color) -> None:
    rect(slide, x, y, w, h, PAPER, LINE, True)
    rect(slide, x, y, 0.055, h, color)
    text(slide, label.upper(), x + 0.20, y + 0.18, w - 0.4, 0.2, 8, color, True)
    text(slide, detail, x + 0.20, y + 0.54, w - 0.4, h - 0.68, 12, INK, True)


def build() -> None:
    required = [LOGO, *(CAPTURE_ROOT / name for name in ("top.png", "overview.png", "facade-side.png", "park-side.png"))]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Ontbrekende presentatiebestanden: " + ", ".join(missing))
    OUTPUT_DIR.mkdir(exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = W
    presentation.slide_height = H
    blank = presentation.slide_layouts[6]

    # 1. Titel
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    slide.shapes.add_picture(str(CAPTURE_ROOT / "overview.png"), Inches(7.58), Inches(0),
                             width=Inches(5.753), height=Inches(7.5))
    rect(slide, 0, 0, 7.78, 7.5, TEAL)
    text(slide, "PARKHAVEN STUDIO / PRESENTATIE V0.2", 0.72, 0.55, 6.2, 0.3, 10, MINT, True)
    text(slide, "Van 3D-model naar\nonderbouwde ontwerpkeuze.", 0.72, 1.35, 6.35, 1.95, 31, WHITE, True)
    text(slide, "Wat de app doet — en hoe de techniek samenwerkt", 0.72, 3.72, 6.15, 0.45, 16, MINT)
    rect(slide, 0.72, 4.38, 1.05, 0.06, ORANGE)
    text(slide, "Rhino • Grasshopper • 3D-bewijs • AI • rapportage", 0.72, 4.72, 6.2, 0.4, 13, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(5.98), Inches(5.40), height=Inches(1.47))
    text(slide, "OVERZICHT + TECHNISCHE ACHTERGROND", 0.72, 6.72, 3.5, 0.2, 8, MINT, True)
    add_notes(slide, scripts[0])

    # 2. Probleem en belofte
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Waarom ParkHaven Studio", "Van losse gereedschappen naar één ontwerpdossier", 2)
    text(slide, "ZONDER APP", 0.78, 1.88, 1.3, 0.2, 8, ORANGE, True)
    fragmented = [("Rhino", BLUE), ("Grasshopper", GOLD), ("Schermafbeeldingen", PURPLE), ("Notities", SOFT_BLUE)]
    for index, (label, color) in enumerate(fragmented):
        y = 2.25 + index * 0.78
        rect(slide, 0.72 + (index % 2) * 1.82, y, 1.55, 0.54, PAPER, color, True)
        text(slide, label, 0.80 + (index % 2) * 1.82, y + 0.16, 1.39, 0.2, 10, color, True,
             align=PP_ALIGN.CENTER)
    arrow(slide, 4.12, 3.60, 5.05, 3.60, MUTED, 2)
    rect(slide, 5.18, 1.98, 3.0, 3.25, TEAL, TEAL, True)
    slide.shapes.add_picture(str(LOGO), Inches(6.20), Inches(2.26), height=Inches(0.88))
    text(slide, "PARKHAVEN\nSTUDIO", 5.58, 3.28, 2.20, 0.75, 21, WHITE, True,
         align=PP_ALIGN.CENTER)
    text(slide, "Eén begeleide workflow", 5.55, 4.45, 2.26, 0.3, 11, MINT,
         align=PP_ALIGN.CENTER)
    arrow(slide, 8.34, 3.60, 9.27, 3.60, GREEN, 2)
    text(slide, "MET APP", 9.45, 1.88, 1.3, 0.2, 8, GREEN, True)
    outcomes = [("Model + versie", BLUE), ("Resultaten", GREEN), ("Beelden", PURPLE), ("Rapport", ORANGE)]
    for index, (label, color) in enumerate(outcomes):
        y = 2.25 + index * 0.78
        rect(slide, 9.38, y, 2.78, 0.54, PAPER, color, True)
        text(slide, label, 9.58, y + 0.16, 2.38, 0.2, 11, INK, True)
        text(slide, "✓", 11.72, y + 0.14, 0.26, 0.22, 12, color, True,
             align=PP_ALIGN.CENTER)
    rect(slide, 0.72, 5.72, 11.44, 0.64, PAPER, LINE, True)
    text(slide, "Iedere conclusie blijft gekoppeld aan dezelfde modelversie, berekening en hetzelfde bewijs.",
         1.00, 5.94, 10.88, 0.24, 13, INK, True, align=PP_ALIGN.CENTER)
    footer_nl(slide)
    add_notes(slide, scripts[1])

    # 3. Gebruikersworkflow
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Gebruikerservaring", "Vijf stappen van model naar besluit", 3)
    labels = [
        ("Importeren", "Bronmodel bewaren", BLUE),
        ("Positioneren", "Eenheden + positie", GOLD),
        ("Analyseren", "Rhino + Grasshopper", GREEN),
        ("Begrijpen", "Viewer + vaste beelden", PURPLE),
        ("Rapporteren", "Advies + PDF", ORANGE),
    ]
    for index, (label, detail, color) in enumerate(labels):
        x = 0.62 + index * 2.45
        step_card(slide, x, index + 1, label, detail, color)
        if index < len(labels) - 1:
            arrow(slide, x + 2.12, 3.43, x + 2.39, 3.43, MUTED, 1.5)
    rect(slide, 0.62, 5.18, 12.08, 0.92, TEAL, TEAL, True)
    text(slide, "De app begeleidt de route; de ontwerper blijft degene die plaatsing, doelen en conclusies goedkeurt.",
         0.92, 5.46, 11.48, 0.3, 13, WHITE, True, align=PP_ALIGN.CENTER)
    footer_nl(slide)
    add_notes(slide, scripts[2])

    # 4. Werkruimte
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Centrale werkruimte", "Vrij onderzoeken én gecontroleerd vergelijken", 4)
    slide.shapes.add_picture(str(CAPTURE_ROOT / "overview.png"), Inches(0.62), Inches(1.72),
                             width=Inches(7.15), height=Inches(4.77))
    rect(slide, 8.02, 1.72, 4.68, 4.77, PAPER, LINE, True)
    text(slide, "3D-VIEWER", 8.34, 2.02, 1.2, 0.2, 8, GREEN, True)
    features = [
        ("Context", "Ontwerp, referenties, parkdoel en luchtfoto"),
        ("Lagen", "Zichtbaarheid per analyse- en modellaag"),
        ("Camera", "Vrij navigeren of vaste kijkrichtingen kiezen"),
        ("Herkenning", "Stabiele kleuren voor Volume A, B en C"),
    ]
    for index, (heading, detail) in enumerate(features):
        y = 2.48 + index * 0.82
        rect(slide, 8.34, y, 0.34, 0.34, (BLUE, GREEN, GOLD, PURPLE)[index],
             (BLUE, GREEN, GOLD, PURPLE)[index], True)
        text(slide, heading, 8.84, y - 0.01, 1.2, 0.24, 11, INK, True)
        text(slide, detail, 10.02, y - 0.01, 2.28, 0.42, 9, MUTED)
    pill(slide, "Volume A", 8.34, 5.92, 1.10, BLUE)
    pill(slide, "Volume B", 9.58, 5.92, 1.10, "F2994A")
    pill(slide, "Volume C", 10.82, 5.92, 1.10, PURPLE)
    footer_nl(slide)
    add_notes(slide, scripts[3])

    # 5. Resultaten
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Wat de app oplevert", "Van kerngetal naar leesbaar ontwerpdossier", 5)
    metric_card(slide, 0.62, 1.70, 3.75, "Bruto vloeroppervlak", "51.868", " m²", GREEN)
    metric_card(slide, 4.55, 1.70, 3.75, "Zicht", "24,5", "%", BLUE)
    metric_card(slide, 8.48, 1.70, 4.22, "Indicatieve kosten", "€ 93,4", " mln", ORANGE)
    domains = [
        ("Zon & klimaat", "Zonuren, UTCI en hittestress als vergelijkende screening", GOLD),
        ("Fundering", "Paalindicatie, belasting en benuttingsgraad", SOFT_BLUE),
        ("Ontwerpadvies", "Meetbare acties, bewijsreferenties en onzekerheden", PURPLE),
        ("Rapport", "Vier beelden, legenda, meetwaarden, methode en beperkingen", GREEN),
    ]
    for index, (heading, detail, color) in enumerate(domains):
        x = 0.62 + (index % 2) * 6.13
        y = 3.10 + (index // 2) * 1.42
        rect(slide, x, y, 5.95, 1.12, PAPER, LINE, True)
        rect(slide, x, y, 0.06, 1.12, color)
        text(slide, heading, x + 0.24, y + 0.20, 1.72, 0.3, 14, INK, True)
        text(slide, detail, x + 2.02, y + 0.18, 3.57, 0.55, 10, MUTED)
    text(slide, "SCREENING ≠ DEFINITIEF ONTWERP", 9.56, 6.19, 3.14, 0.22, 8, ORANGE, True,
         align=PP_ALIGN.RIGHT)
    footer_nl(slide)
    add_notes(slide, scripts[4])

    # 6. Architectuur
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Architectuur", "Zes componenten, één lokale workflow", 6)
    component(slide, 0.62, 1.76, 2.25, 1.25, "Desktop", "WPF + WebView2\nvenster en levenscyclus", BLUE)
    component(slide, 3.20, 1.76, 2.25, 1.25, "Backend", "FastAPI\nprojecten, runs en API", GREEN)
    component(slide, 5.78, 1.76, 2.25, 1.25, "Rekenkern", "Rhino-plug-in +\nGrasshopper-definitie", GOLD)
    component(slide, 8.36, 1.76, 2.25, 1.25, "Viewer", "Three.js\nscene en bewijs", PURPLE)
    component(slide, 10.94, 1.76, 1.76, 1.25, "Rapport", "HTML → PDF", ORANGE)
    for x in (2.87, 5.45, 8.03, 10.61):
        arrow(slide, x, 2.38, x + 0.27, 2.38, MUTED, 1.4)
    rect(slide, 0.62, 3.50, 7.41, 2.23, TEAL, TEAL, True)
    text(slide, "LOKALE DATAROUTE", 0.92, 3.80, 1.8, 0.22, 8, MINT, True)
    data_steps = [("PROJECT", "bron + werkmodel"), ("RUN", "job + kopieën"),
                  ("DATA", "meetwaarden + scene"), ("OUTPUT", "viewer + rapport")]
    for index, (heading, detail) in enumerate(data_steps):
        x = 0.92 + index * 1.70
        text(slide, heading, x, 4.32, 1.25, 0.22, 10, WHITE, True)
        text(slide, detail, x, 4.70, 1.36, 0.38, 9, MINT)
        if index < len(data_steps) - 1:
            arrow(slide, x + 1.24, 4.45, x + 1.54, 4.45, "9DBDB5", 1.3)
    rect(slide, 8.36, 3.50, 4.34, 2.23, PAPER, LINE, True)
    text(slide, "OPTIONELE EXTERNE ROUTE", 8.66, 3.80, 2.2, 0.22, 8, ORANGE, True)
    text(slide, "Begrensd tekstpakket", 8.66, 4.28, 1.98, 0.28, 11, INK, True)
    arrow(slide, 10.72, 4.43, 11.08, 4.43, ORANGE, 1.8)
    text(slide, "AI-endpoint", 11.18, 4.28, 1.18, 0.28, 12, INK, True)
    text(slide, "Geen afbeeldingen\nof ruwe geometrie", 8.66, 4.91, 3.40, 0.45, 9, MUTED)
    footer_nl(slide)
    add_notes(slide, scripts[5])

    # 7. Run-isolatie
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Technische verdieping I", "Iedere analyse is een geïsoleerde, herleidbare run", 7)
    component(slide, 0.62, 1.78, 2.22, 1.16, "Bron", "original-design.3dm\nongewijzigd", BLUE)
    arrow(slide, 2.84, 2.36, 3.30, 2.36, MUTED, 1.8)
    component(slide, 3.34, 1.78, 2.22, 1.16, "Werkmodel", "parkhaven-design.3dm\nplaatsing goedgekeurd", GOLD)
    arrow(slide, 5.56, 2.36, 6.02, 2.36, MUTED, 1.8)
    rect(slide, 6.06, 1.62, 6.64, 3.92, PAPER, LINE, True)
    text(slide, "RUN / UUID", 6.36, 1.92, 1.5, 0.22, 9, GREEN, True)
    files = [
        ("job.json", "status, versies, timing"),
        ("design.3dm", "runlokale modelkopie"),
        ("ParkHaven_Analysis.ghx", "vaste rekendefinitie"),
        ("weather/", "gepinde weersbron"),
        ("data/latest.json", "berekende meetwaarden"),
        ("data/scene.json", "3D-scene voor viewer"),
        ("data/view-evidence.json", "camera's, lagen en hashes"),
    ]
    for index, (name, detail) in enumerate(files):
        y = 2.36 + index * 0.39
        text(slide, name, 6.40, y, 2.58, 0.20, 9, INK, True)
        text(slide, detail, 9.08, y, 2.98, 0.20, 8, MUTED)
    rect(slide, 0.62, 3.56, 4.94, 1.98, TEAL, TEAL, True)
    text(slide, "PROVENANCE", 0.92, 3.88, 1.35, 0.22, 8, MINT, True)
    text(slide, "SHA-256 koppelt bron, definitie, weer,\nscene, resultaten en beelden.",
         0.92, 4.34, 4.14, 0.72, 15, WHITE, True)
    rect(slide, 6.06, 5.86, 6.64, 0.42, TEAL, TEAL, True)
    text(slide, "Atomisch publiceren • geen stille overschrijving • verouderd bewijs wordt geweigerd",
         6.28, 5.99, 6.20, 0.17, 8, WHITE, True, align=PP_ALIGN.CENTER)
    footer_nl(slide)
    add_notes(slide, scripts[6])

    # 8. Evidence en AI-grens
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, BG)
    title(slide, "Technische verdieping II", "Van scene naar bewijs — met een harde datagrens", 8)
    pipeline = [
        (0.62, "SCENE", "mesh + kleuren", BLUE),
        (2.72, "4 BEELDEN", "vaste camera's", PURPLE),
        (4.82, "BEWIJS", "meetwaarden + volumes", GREEN),
        (6.92, "AI", "tekstpakket", ORANGE),
        (9.02, "VALIDATIE", "beleid + IDs", GOLD),
        (11.12, "RAPPORT", "actueel bewijs", SOFT_BLUE),
    ]
    for index, (x, heading, detail, color) in enumerate(pipeline):
        rect(slide, x, 2.00, 1.58, 1.18, PAPER, LINE, True)
        rect(slide, x, 2.00, 1.58, 0.07, color)
        text(slide, heading, x + 0.14, 2.30, 1.30, 0.22, 9, color, True,
             align=PP_ALIGN.CENTER)
        text(slide, detail, x + 0.14, 2.72, 1.30, 0.22, 8, MUTED,
             align=PP_ALIGN.CENTER)
        if index < len(pipeline) - 1:
            arrow(slide, x + 1.58, 2.59, x + 2.04, 2.59, MUTED, 1.3)
    rect(slide, 0.62, 3.80, 5.78, 1.84, TEAL, TEAL, True)
    text(slide, "BLIJFT LOKAAL", 0.92, 4.10, 1.45, 0.22, 8, MINT, True)
    text(slide, "3D-model • screenshots • bronbestanden", 0.92, 4.56, 4.95, 0.34, 15, WHITE, True)
    text(slide, "Alleen een begrensde tekstsamenvatting verlaat de machine.",
         0.92, 5.08, 4.95, 0.3, 9, MINT)
    rect(slide, 6.68, 3.80, 6.02, 1.84, PAPER, LINE, True)
    text(slide, "SERVERCONTROLE", 6.98, 4.10, 1.65, 0.22, 8, ORANGE, True)
    checks = ["bekende volumes", "bekende meetwaarden", "huidige beleidsversie", "toegestane ingreep"]
    for index, check in enumerate(checks):
        x = 6.98 + (index % 2) * 2.55
        y = 4.58 + (index // 2) * 0.48
        text(slide, "✓", x, y, 0.22, 0.22, 10, GREEN, True)
        text(slide, check, x + 0.30, y, 2.05, 0.22, 9, INK, True)
    text(slide, "Verschuivingen zijn verkennend → plot- en tunnelcontrole blijft vereist",
         6.98, 5.74, 5.40, 0.25, 8, ORANGE, True)
    footer_nl(slide)
    add_notes(slide, scripts[7])

    # 9. Status en roadmap
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, TEAL)
    text(slide, "WAAR WE NU STAAN / V0.2", 0.72, 0.52, 4.5, 0.3, 10, MINT, True)
    text(slide, "De kernworkflow staat.", 0.72, 1.08, 6.2, 0.62, 30, WHITE, True)
    text(slide, "Van lokaal model tot controleerbaar rapport", 0.72, 1.87, 5.9, 0.35, 15, MINT)
    capabilities = [
        ("Project", "import + bronbehoud"), ("Plaatsing", "preview + goedkeuring"),
        ("Analyse", "Rhino + Grasshopper"), ("Bewijs", "viewer + 4 beelden"),
        ("Advies", "begrensd + gevalideerd"), ("Rapport", "Engels + PDF"),
    ]
    for index, (heading, detail) in enumerate(capabilities):
        x = 0.72 + (index % 3) * 2.95
        y = 2.66 + (index // 3) * 1.10
        rect(slide, x, y, 2.66, 0.82, "214849", "496766", True)
        text(slide, heading, x + 0.18, y + 0.16, 1.05, 0.22, 11, WHITE, True)
        text(slide, detail, x + 1.13, y + 0.16, 1.30, 0.38, 8, MINT)
    rect(slide, 9.82, 2.66, 2.78, 1.92, ORANGE, ORANGE, True)
    text(slide, "VOLGENDE FASE", 10.12, 2.98, 2.18, 0.22, 9, WHITE, True,
         align=PP_ALIGN.CENTER)
    text(slide, "Betrouwbaarheid\n+ praktijktest", 10.12, 3.48, 2.18, 0.62, 17, WHITE, True,
         align=PP_ALIGN.CENTER)
    rect(slide, 0.72, 5.14, 11.88, 0.02, "496766")
    roadmap = ["Geverifieerde locatiegrenzen", "Gekalibreerde doelen en normen",
               "Ondertekende distributie", "Schone testmachines"]
    for index, item in enumerate(roadmap):
        x = 0.72 + index * 2.95
        text(slide, f"0{index + 1}", x, 5.55, 0.34, 0.24, 9, ORANGE, True)
        text(slide, item, x + 0.44, 5.50, 2.20, 0.55, 11, WHITE, True)
    text(slide, "OVERZICHTELIJK  •  MEETBAAR  •  HERLEIDBAAR", 0.72, 6.62, 6.4, 0.3, 10, MINT, True)
    slide.shapes.add_picture(str(LOGO), Inches(11.75), Inches(5.72), height=Inches(1.13))
    add_notes(slide, scripts[8])

    presentation.core_properties.title = "ParkHaven Studio presentatie v0.2"
    presentation.core_properties.subject = "Nederlandse productintroductie en technische achtergrond"
    presentation.core_properties.author = "ParkHaven Studio"
    presentation.save(OUTPUT)

    sections = [
        "Introductie", "Waarom", "Gebruikersworkflow", "Werkruimte", "Resultaten",
        "Architectuur", "Run-isolatie", "Bewijs en AI", "Status en roadmap",
    ]
    lines = ["# ParkHaven Studio — presentatiescript v0.2", ""]
    for index, (section, script) in enumerate(zip(sections, scripts), start=1):
        lines.extend((f"## Dia {index}: {section}", script, ""))
    SCRIPT_OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Gemaakt: {OUTPUT}")
    print(f"Gemaakt: {SCRIPT_OUTPUT}")


if __name__ == "__main__":
    build()
